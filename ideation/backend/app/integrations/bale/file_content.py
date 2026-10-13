import re
import subprocess
import tempfile
import unicodedata
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

from app.core.config import settings

ALLOWED_BALE_EXTENSIONS = {".pdf", ".doc", ".docx", ".txt"}
MIME_TYPES_BY_EXTENSION = {
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".doc": {"application/msword", "application/octet-stream"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
        "application/zip",
    },
    ".txt": {"text/plain", "application/octet-stream"},
}
DEFAULT_MIME_BY_EXTENSION = {
    ".pdf": "application/pdf",
    ".doc": "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
}
MAX_CLASSIFICATION_CHARS = 100_000
MAX_PDF_CLASSIFICATION_PAGES = 200
MAX_DOCX_EXPANDED_BYTES = 100 * 1024 * 1024


class BaleContentValidationError(ValueError):
    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


@dataclass(frozen=True)
class PreparedBaleDocument:
    filename: str
    extension: str
    mime_type: str
    content: bytes
    classification_text: str


def _safe_filename(value: str | None) -> tuple[str, str]:
    if not value:
        raise BaleContentValidationError("نام فایل مشخص نیست؛ لطفاً فایل را دوباره ارسال کنید.")
    filename = re.split(r"[\\/]", unicodedata.normalize("NFKC", value))[-1]
    filename = "".join(char for char in filename if char.isprintable()).strip()
    if not filename or len(filename) > 240:
        raise BaleContentValidationError("نام فایل معتبر نیست یا بیش از حد طولانی است.")
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_BALE_EXTENSIONS:
        raise BaleContentValidationError(
            "فقط فایل‌های PDF، Word با پسوند DOC یا DOCX و فایل TXT پذیرفته می‌شوند."
        )
    return filename, extension


def _validate_mime(extension: str, mime_type: str | None) -> str:
    normalized = (mime_type or "").split(";", 1)[0].strip().lower()
    if normalized and normalized not in MIME_TYPES_BY_EXTENSION[extension]:
        raise BaleContentValidationError("نوع واقعی فایل با پسوند آن سازگار نیست.")
    return normalized or DEFAULT_MIME_BY_EXTENSION[extension]


def _validate_signature(extension: str, content: bytes) -> None:
    if not content:
        raise BaleContentValidationError("فایل ارسالی خالی است.")
    if extension == ".pdf" and not content.startswith(b"%PDF-"):
        raise BaleContentValidationError("ساختار فایل PDF معتبر نیست.")
    if extension == ".doc" and not content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        raise BaleContentValidationError("ساختار فایل DOC معتبر نیست.")
    if extension == ".docx":
        try:
            with ZipFile(BytesIO(content)) as archive:
                names = set(archive.namelist())
                if (
                    len(names) > 10_000
                    or sum(item.file_size for item in archive.infolist())
                    > MAX_DOCX_EXPANDED_BYTES
                ):
                    raise BaleContentValidationError("ساختار فایل DOCX بیش از حد بزرگ است.")
                if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise BaleContentValidationError("ساختار فایل DOCX معتبر نیست.")
        except BadZipFile as exc:
            raise BaleContentValidationError("ساختار فایل DOCX معتبر نیست.") from exc
    if extension == ".txt" and b"\x00" in content[:8192]:
        raise BaleContentValidationError("فایل TXT حاوی داده باینری است.")


def _extract_txt(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1256"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise BaleContentValidationError("رمزگذاری فایل TXT قابل تشخیص نیست.")


def _extract_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted and not reader.decrypt(""):
            raise BaleContentValidationError("فایل PDF رمزگذاری‌شده قابل پردازش نیست.")
        parts: list[str] = []
        total_chars = 0
        for index, page in enumerate(reader.pages):
            if index >= MAX_PDF_CLASSIFICATION_PAGES:
                break
            value = page.extract_text() or ""
            parts.append(value)
            total_chars += len(value)
            if total_chars >= MAX_CLASSIFICATION_CHARS:
                break
        return "\n".join(parts)
    except BaleContentValidationError:
        raise
    except Exception as exc:
        raise BaleContentValidationError("متن فایل PDF قابل استخراج نیست.") from exc


def _extract_docx(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise BaleContentValidationError("متن فایل DOCX قابل استخراج نیست.") from exc


def _extract_doc(content: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".doc") as temporary:
        temporary.write(content)
        temporary.flush()
        try:
            result = subprocess.run(
                ["antiword", "-m", "UTF-8.txt", temporary.name],
                check=False,
                capture_output=True,
                timeout=15,
            )
        except (FileNotFoundError, subprocess.TimeoutExpired):
            return ""
    if result.returncode != 0:
        return ""
    return result.stdout.decode("utf-8", errors="replace")


def validate_bale_document_metadata(
    *,
    filename: str | None,
    mime_type: str | None,
    reported_size: int | None,
) -> tuple[str, str, str]:
    safe_filename, extension = _safe_filename(filename)
    safe_mime = _validate_mime(extension, mime_type)
    if reported_size is not None and reported_size < 0:
        raise BaleContentValidationError("اندازه اعلام‌شده فایل معتبر نیست.")
    if reported_size is not None and reported_size > settings.bale_file_upload_max_bytes:
        raise BaleContentValidationError(
            f"حجم فایل بیشتر از سقف مجاز {settings.bale_file_upload_max_mb} مگابایت است."
        )
    return safe_filename, extension, safe_mime


def prepare_bale_document(
    *,
    filename: str | None,
    mime_type: str | None,
    reported_size: int | None,
    content: bytes,
    caption: str | None,
) -> PreparedBaleDocument:
    safe_filename, extension, safe_mime = validate_bale_document_metadata(
        filename=filename,
        mime_type=mime_type,
        reported_size=reported_size,
    )
    limit = settings.bale_file_upload_max_bytes
    if len(content) > limit:
        raise BaleContentValidationError(
            f"حجم فایل بیشتر از سقف مجاز {settings.bale_file_upload_max_mb} مگابایت است."
        )
    _validate_signature(extension, content)

    if extension == ".txt":
        extracted = _extract_txt(content)
    elif extension == ".pdf":
        extracted = _extract_pdf(content)
    elif extension == ".docx":
        extracted = _extract_docx(content)
    else:
        extracted = _extract_doc(content)
    classification_text = "\n".join(
        part.strip()
        for part in (caption or "", safe_filename, extracted)
        if part and part.strip()
    )[:MAX_CLASSIFICATION_CHARS]
    return PreparedBaleDocument(
        filename=safe_filename,
        extension=extension,
        mime_type=safe_mime,
        content=content,
        classification_text=classification_text,
    )
